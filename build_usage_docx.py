# -*- coding: utf-8 -*-
"""把《使用说明.md》（含截图）生成带图、按节分页的 Word《使用说明.docx》。

- 标题 → Heading（微软雅黑），每个 “## ” 章节前插入分页（逐页对照 App 多页报告）
- 图片行 ![图注](docs/img/x.png) → 居中插图 + 图注
- GFM 表 → Word 表格（表头加粗底纹）
- 加粗/行内代码、引用、项目符号、编号段落
"""
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).resolve().parent
SRC = BASE / "使用说明.md"
OUT = BASE / "使用说明.docx"

HEAD_RE = re.compile(r"^(#{1,6})\s+(.*)$")
IMG_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
SEP_RE = re.compile(r"^\s*\|?[\s:|-]*-[\s:|-]*\|?\s*$")
TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def set_ea(run, ea="宋体"):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), ea)


def add_runs(p, text, size=None, ea="宋体", bold=False):
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            _plain(p, text[pos:m.start()], size, ea, bold)
        tok = m.group(0)
        if tok.startswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
            if size: r.font.size = size
            set_ea(r, ea)
        else:
            r = p.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt((size.pt - 1) if size else 9)
            rf = r._element.get_or_add_rPr().find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts"); r._element.get_or_add_rPr().append(rf)
            rf.set(qn("w:ascii"), "Consolas"); rf.set(qn("w:hAnsi"), "Consolas")
        pos = m.end()
    if pos < len(text):
        _plain(p, text[pos:], size, ea, bold)


def _plain(p, text, size, ea, bold):
    if not text:
        return
    r = p.add_run(text)
    if bold: r.bold = True
    if size: r.font.size = size
    set_ea(r, ea)


def shade(cell, hexcolor="D9E2F3"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def split_cells(line):
    line = line.strip()
    if line.startswith("|"): line = line[1:]
    if line.endswith("|"): line = line[:-1]
    return [c.strip() for c in line.split("|")]


def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            p = cell.paragraphs[0]
            if ri == 0:
                add_runs(p, val, size=Pt(9), ea="微软雅黑")
                for r in p.runs: r.bold = True
                shade(cell)
            else:
                add_runs(p, val, size=Pt(9))
    doc.add_paragraph()


def add_image(doc, path, alt):
    if os.path.exists(path):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Inches(6.2))
        cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(alt); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x59, 0x63, 0x74); r.italic = True
        set_ea(r)
    else:
        p = doc.add_paragraph()
        r = p.add_run(f"[图片缺失：{path}]"); r.italic = True


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1); s.bottom_margin = Inches(1)
        s.left_margin = Inches(1); s.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    lines = SRC.read_text(encoding="utf-8").splitlines()
    in_code = None
    buf = []
    first = True
    img_count = 0

    def flush_code():
        nonlocal buf
        if not buf: return
        for ln in buf:
            p = doc.add_paragraph()
            r = p.add_run(ln); r.font.name = "Consolas"; r.font.size = Pt(9); set_ea(r)
        buf = []

    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i].rstrip()
        st = raw.strip()

        if st.startswith("```"):
            if in_code is None:
                in_code = True
            else:
                flush_code(); in_code = None
            i += 1
            continue
        if in_code:
            buf.append(raw); i += 1; continue
        if not st:
            i += 1; continue

        # 图片
        m = IMG_RE.match(st)
        if m:
            add_image(doc, str(BASE / m.group(2)), m.group(1) or "插图")
            img_count += 1
            i += 1
            continue

        # 标题
        m = HEAD_RE.match(st)
        if m:
            lvl = min(len(m.group(1)), 6)
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2)).replace("`", "")
            if lvl <= 2 and not first:
                doc.add_page_break()          # 每节另起一页
            first = False
            h = doc.add_heading(level=lvl)
            h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)
            r = h.add_run(text)
            r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F); r.font.name = "微软雅黑"
            set_ea(r, "微软雅黑")
            i += 1
            continue

        # 表格
        if st.startswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i].strip()); i += 1
            if len(block) >= 2 and SEP_RE.match(block[1]) and "-" in block[1]:
                header = split_cells(block[0])
                data = [split_cells(b) for b in block[2:]]
                add_table(doc, [header] + data)
            else:
                add_table(doc, [split_cells(b) for b in block])
            continue

        # 引用
        if st.startswith(">"):
            text = st.lstrip(">").strip()
            if text.startswith("!"):
                m2 = IMG_RE.match(text)
                if m2:
                    add_image(doc, str(BASE / m2.group(2)), m2.group(1)); img_count += 1
                    i += 1; continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            add_runs(p, text)
            i += 1
            continue

        # 项目符号
        if st.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, st[2:])
            i += 1
            continue

        # 编号
        if re.match(r"^\d+[\.、]", st):
            p = doc.add_paragraph()
            add_runs(p, st)
            i += 1
            continue

        # 分隔线
        if re.match(r"^-{3,}$", st):
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, st)
        i += 1

    flush_code()
    doc.save(str(OUT))
    print("已生成:", OUT)
    print("段落:", len(doc.paragraphs), "表格:", len(doc.tables), "图片:", img_count)


if __name__ == "__main__":
    main()
